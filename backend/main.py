import io
import os
import sys
import asyncio
from pathlib import Path

# Ensure the backend directory is in sys.path so sibling imports work everywhere
backend_dir = Path(__file__).resolve().parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from fastapi import FastAPI, File, UploadFile, HTTPException, status, Request
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
import pypdf
import docx
from typing import Optional, Literal, List
from pydantic import BaseModel, Field, field_validator
from dotenv import load_dotenv

from analyze_contract import (
    analyze_contract_text,
    ContractAnalysisResult,
    ClauseAnalysis,
    generate_negotiation_draft,
    simulate_scenario,
)
from generate_report import generate_report_pdf
from generate_docx import generate_docx_track_changes

# Maximum allowed file upload size: 15 MB
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024


# ==============================================================================
# Pydantic Schemas for Feature 1 (Negotiation Drafts)
# ==============================================================================

class NegotiationDraftRequest(BaseModel):
    """Payload for generating diplomatic and firm negotiation counter-proposals."""
    clause_title: str = Field(..., min_length=2, max_length=200, description="Clause title")
    original_text: str = Field(..., min_length=5, max_length=15000, description="Original contract clause")
    suggested_redline: str = Field(..., min_length=5, max_length=15000, description="Counter-proposal redline")
    counterparty_name: Optional[str] = Field(default="Client / Landlord", max_length=150)
    user_role: Optional[str] = Field(default="Contractor / Tenant", max_length=150)

    @field_validator("clause_title", "original_text", "suggested_redline", mode="before")
    @classmethod
    def sanitize_strings(cls, v: str) -> str:
        """Strip whitespace and enforce string hygiene."""
        if isinstance(v, str):
            return v.strip()
        return v


class NegotiationDraftResponse(BaseModel):
    """Multi-channel, multi-tone structured negotiation draft outputs."""
    email_subject: str
    email_diplomatic: str
    email_firm: str
    chat_diplomatic: str
    chat_firm: str


# ==============================================================================
# Pydantic Schemas for Feature 3 (What-If Scenario Simulator)
# ==============================================================================

class ScenarioRequest(BaseModel):
    """Payload for evaluating a hypothetical user scenario against an analyzed contract."""
    scenario_query: str = Field(
        ...,
        min_length=3,
        max_length=500,
        description="Hypothetical situation or question to simulate against contract terms"
    )
    contract_data: ContractAnalysisResult

    @field_validator("scenario_query", mode="before")
    @classmethod
    def sanitize_scenario_query(cls, v: str) -> str:
        """Strip leading/trailing whitespace and validate length."""
        if isinstance(v, str):
            cleaned = v.strip()
            if len(cleaned) < 3:
                raise ValueError("Scenario query must be at least 3 characters after trimming.")
            return cleaned
        return v


class ScenarioSimulationResponse(BaseModel):
    """Deterministic simulation results including badge, quote, and advice."""
    scenario_query: str
    verdict_badge: Literal["SAFE", "AT_RISK", "SEVERE_PENALTY", "UNADDRESSED"]
    direct_consequence: str
    governing_clause_title: str
    governing_clause_quote: str
    recommended_action: str


# ==============================================================================
# Pydantic Schemas for Feature 4 (Track-Changes Word Export)
# ==============================================================================

class DocxExportRequest(BaseModel):
    """Payload for generating a Word document (.docx) with Track-Changes styling."""
    document_title: str = Field(default="Contract Agreement", max_length=250)
    original_full_text: Optional[str] = ""
    clauses: List[ClauseAnalysis]


# Load environment variables
load_dotenv()
load_dotenv(Path(__file__).parent.parent / ".env")
load_dotenv(Path(__file__).parent / ".env")

# Initialize Slowapi Rate Limiter
limiter = Limiter(key_func=get_remote_address)

app = FastAPI(
    title="LegalDoc.AI Contract Analysis API",
    description="Deterministic contract risk assessment, native language translation, and redline negotiation engine",
    version="1.0.0",
)

app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

# Configure CORS to allow frontend connections
origins = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    "http://localhost:3001",
    "http://127.0.0.1:3001",
    "http://localhost:5173",
    "http://127.0.0.1:5173",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Allow all origins for seamless local dev
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def extract_text_from_file(filename: str, content: bytes) -> str:
    """
    Extracts plain text from PDF, DOCX, and TXT file bytes.
    """
    ext = Path(filename).suffix.lower()

    if ext == ".txt":
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")

    elif ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_pages.append(text.strip())
            return "\n\n".join(extracted_pages)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse PDF document: {str(e)}",
            )

    elif ext in [".docx", ".doc"]:
        try:
            doc = docx.Document(io.BytesIO(content))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            return "\n\n".join(full_text)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse Word (.docx) document: {str(e)}",
            )

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Please upload a .pdf, .docx, or .txt file.",
        )


@app.get("/")
async def root():
    return {
        "service": "LegalDoc.AI Contract Analysis API",
        "docs": "/docs",
        "health": "/api/health",
        "version": "1.0.0"
    }


@app.get("/api/health")
async def health_check():
    """
    Health check endpoint for the backend service.
    """
    return {
        "status": "healthy",
        "service": "LegalDoc.AI Backend",
        "gemini_configured": bool(os.getenv("GEMINI_API_KEY")),
    }


@app.post("/api/analyze", response_model=ContractAnalysisResult, response_class=JSONResponse)
@limiter.limit("10/minute")
async def analyze_contract(request: Request, file: UploadFile = File(...)):
    """
    Analyzes an uploaded contract document (.pdf, .docx, .txt) and returns
    structured risk evaluation, redlines, and plain language translation
    in the document's native language with preserved UTF-8 encoding.
    Protected by 10 req/min rate limit and 15MB file size limit.
    """
    filename = file.filename or "unknown_contract.txt"
    ext = Path(filename).suffix.lower()

    if ext not in [".pdf", ".docx", ".doc", ".txt"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .txt",
        )

    try:
        content = await file.read()
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Error reading uploaded file: {str(e)}",
        )

    if not content or len(content) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file is empty.",
        )

    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Uploaded file exceeds maximum limit of 15MB (received {len(content)} bytes).",
        )

    extracted_text = extract_text_from_file(filename, content)

    if not extracted_text or len(extracted_text.strip()) < 20:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Could not extract sufficient text from the uploaded document. Please check the file contents.",
        )

    try:
        result = await asyncio.to_thread(analyze_contract_text, extracted_text)
        return JSONResponse(
            content=result.model_dump(),
            media_type="application/json; charset=utf-8"
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"AI Contract Analysis failed: {str(e)}",
        )


@app.post("/api/export-pdf")
async def export_pdf_report(analysis: ContractAnalysisResult):
    """
    Generates an executive, styled PDF report of the contract risk analysis.
    """
    try:
        pdf_buffer = await asyncio.to_thread(generate_report_pdf, analysis)
        doc_slug = "Contract"
        if analysis.document_title:
            doc_slug = "".join(c for c in analysis.document_title if c.isalnum() or c in (" ", "_", "-")).strip()
            doc_slug = doc_slug.replace(" ", "_")[:30] or "Contract"
        filename = f"LegalDoc_{doc_slug}_Report.pdf"

        return StreamingResponse(
            pdf_buffer,
            media_type="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition",
            },
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate PDF report: {str(e)}",
        )


@app.post("/api/export-docx")
async def export_docx_amended_agreement(req: DocxExportRequest):
    """
    Generates a Word (.docx) amended agreement with Track-Changes styling:
    - Struck-through deleted terms in red
    - Underlined, highlighted counter-proposals in green
    - Plain-English rationale and negotiation annotations
    """
    try:
        docx_buffer = await asyncio.to_thread(
            generate_docx_track_changes,
            document_title=req.document_title,
            clauses=req.clauses,
            original_full_text=req.original_full_text or "",
        )
        doc_slug = "Contract"
        if req.document_title:
            doc_slug = "".join(c for c in req.document_title if c.isalnum() or c in (" ", "_", "-")).strip()
            doc_slug = doc_slug.replace(" ", "_")[:30] or "Contract"
        filename = f"LegalDoc_{doc_slug}_Amended_Redline.docx"

        return StreamingResponse(
            docx_buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition",
            },
        )
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate Word (.docx) document: {str(e)}",
        )


@app.post("/api/negotiate-draft", response_model=NegotiationDraftResponse)
async def draft_negotiation_message(req: NegotiationDraftRequest):
    """
    Generates ready-to-send negotiation email and WhatsApp/Slack drafts
    (in both 'Diplomatic' and 'Firm' tones) proposing a balanced counter-proposal redline.
    """
    try:
        response = await asyncio.to_thread(generate_negotiation_draft, req)
        return response
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate negotiation draft: {str(e)}",
        )



@app.post("/api/simulate-scenario", response_model=ScenarioSimulationResponse)
@limiter.limit("10/minute")
async def simulate_contract_scenario(request: Request, req: ScenarioRequest):
    """
    Simulates a 'What If?' hypothetical scenario against the analyzed contract,
    returning a grounded consequence analysis with direct clause citations and verdict badge.
    Protected by 10 req/min rate limit.
    """
    try:
        response = await asyncio.to_thread(simulate_scenario, req)
        return response
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to simulate scenario: {str(e)}",
        )


if __name__ == "__main__":

    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
